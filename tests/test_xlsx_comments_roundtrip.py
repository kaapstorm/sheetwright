import tempfile
from contextlib import contextmanager
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment as XComment
from testsweet import test

from sheetwright.xlsx.reader import read_xlsx
from sheetwright.xlsx.writer import write_xlsx
from sheetwright.security import SecurityLimits


@contextmanager
def _tmp_path():
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


def _wb_with_comment(path: Path, addr: str, author: str, text: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    ws[addr] = 1
    ws[addr].comment = XComment(text, author)
    wb.save(path)


@test
def reader_picks_up_comment():
    with _tmp_path() as tmp_path:
        src = tmp_path / 'in.xlsx'
        _wb_with_comment(src, 'B2', 'Alice', 'Check this')
        wb = read_xlsx(src, limits=SecurityLimits.defaults())
        s = wb.sheet('S')
        assert 'B2' in s.comments
        assert s.comments['B2'].author == 'Alice'
        assert s.comments['B2'].text == 'Check this'


@test
def comment_round_trips_through_writer():
    with _tmp_path() as tmp_path:
        src = tmp_path / 'in.xlsx'
        _wb_with_comment(src, 'B2', 'Alice', 'Check this')
        out = tmp_path / 'out.xlsx'
        write_xlsx(read_xlsx(src, limits=SecurityLimits.defaults()), out)
        s = read_xlsx(out, limits=SecurityLimits.defaults()).sheet('S')
        assert s.comments['B2'].text == 'Check this'
        assert s.comments['B2'].author == 'Alice'
