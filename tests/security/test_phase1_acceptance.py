"""Phase 1 security envelope — end-to-end acceptance smoke test.

Exercises the full security surface in one place so integration regressions
(individual unit tests pass but the envelope has a hole) are caught early.
Intentionally redundant with tests/security/test_mcp_path_traversal.py and
tests/security/test_safe_load_workbook.py.
"""

from __future__ import annotations

import dataclasses
import tempfile
from contextlib import contextmanager
from pathlib import Path
from typing import Any

import openpyxl
from testsweet import catch_exceptions, test

from sheetwright.exceptions import XlsxTooLargeError
from sheetwright.mcp.errors import MCPError
from sheetwright.mcp.server import do_build, do_diff, do_init, do_test
from sheetwright.security import SecurityLimits
from sheetwright.xlsx.safe_load import safe_load_workbook
from tests.fixtures.workbooks import write_simple_xlsx


@contextmanager
def _imported_project():
    """Build a small imported+built project for integration testing."""
    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td)
        src = tmp_path / 'in.xlsx'
        write_simple_xlsx(src)
        p = tmp_path / 'proj'
        p.mkdir()
        (p / 'sheetwright.toml').write_text(
            '[project]\nname = "in"\n[build]\ncalc_engine = "libreoffice"\n'
        )
        (p / 'workbook.toml').write_text(
            '[workbook]\nname = "in"\nsheets = []\n'
        )
        (p / 'sheets').mkdir()
        (p / 'data').mkdir()
        (p / 'tests').mkdir()
        from click.testing import CliRunner

        from sheetwright.cli import main

        CliRunner().invoke(main, ['import', str(src), '--project', str(p)])
        CliRunner().invoke(main, ['build', '--project', str(p)])
        yield tmp_path, p


def _limits(**kwargs: Any) -> SecurityLimits:
    return dataclasses.replace(SecurityLimits.defaults(), **kwargs)


# ---------------------------------------------------------------------------
# do_build — out_path containment
# ---------------------------------------------------------------------------


@test
def acceptance_build_out_of_bounds_raises_path_outside_project():
    with _imported_project() as (_tmp, p):
        with catch_exceptions() as excs:
            do_build(project=str(p), out_path='../escape.xlsx')
        assert excs and isinstance(excs[0], MCPError)
        assert excs[0].code == 'path_outside_project'


@test
def acceptance_build_in_bounds_succeeds():
    with _imported_project() as (_tmp, p):
        (p / 'build').mkdir(exist_ok=True)
        out = do_build(
            project=str(p), out_path=str(p / 'build' / 'custom.xlsx')
        )
        assert out['ok'] is True


# ---------------------------------------------------------------------------
# do_diff — vs path containment
# ---------------------------------------------------------------------------


@test
def acceptance_diff_out_of_bounds_raises_path_outside_project():
    with _imported_project() as (_tmp, p):
        with catch_exceptions() as excs:
            do_diff(project=str(p), vs='xlsx:/etc/passwd')
        assert excs and isinstance(excs[0], MCPError)
        assert excs[0].code == 'path_outside_project'


@test
def acceptance_diff_in_bounds_does_not_raise_path_outside_project():
    """Containment passes; the call may fail for an unrelated reason."""
    with _imported_project() as (_tmp, p):
        with catch_exceptions() as excs:
            do_diff(project=str(p), vs=f'xlsx:{p}/build/in.xlsx')
        if excs:
            assert isinstance(excs[0], MCPError)
            assert excs[0].code != 'path_outside_project'


# ---------------------------------------------------------------------------
# do_test — targets containment
# ---------------------------------------------------------------------------


@test
def acceptance_test_out_of_bounds_raises_path_outside_project():
    with _imported_project() as (_tmp, p):
        with catch_exceptions() as excs:
            do_test(project=str(p), targets=['../../etc/x.py'])
        assert excs and isinstance(excs[0], MCPError)
        assert excs[0].code == 'path_outside_project'


@test
def acceptance_test_in_bounds_does_not_raise_path_outside_project():
    """Path-containment passes; missing file is a different error (or no error)."""
    with _imported_project() as (_tmp, p):
        (p / 'tests' / 'test_real.py').write_text(
            'from testsweet import test\n\n@test\ndef placeholder(): pass\n'
        )
        with catch_exceptions() as excs:
            do_test(project=str(p), targets=['tests/test_real.py'])
        if excs:
            assert isinstance(excs[0], MCPError)
            assert excs[0].code != 'path_outside_project'


# ---------------------------------------------------------------------------
# do_init — path containment
# ---------------------------------------------------------------------------


@test
def acceptance_init_out_of_bounds_raises_path_outside_project():
    with catch_exceptions() as excs:
        do_init(path='../escape')
    assert excs and isinstance(excs[0], MCPError)
    assert excs[0].code == 'path_outside_project'


@test
def acceptance_init_in_bounds_succeeds():
    with tempfile.TemporaryDirectory() as td:
        target = Path(td) / 'fresh'
        out = do_init(path=str(target))
        assert out['ok'] is True


# ---------------------------------------------------------------------------
# Resource-budget (size cap) end-to-end
# ---------------------------------------------------------------------------


@test
def acceptance_size_cap_rejects_oversized_xlsx():
    """safe_load_workbook raises XlsxTooLargeError when uncompressed size
    exceeds a tight limit, and loads cleanly under the default limit."""
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / 'tiny.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'hello'
        ws['B2'] = 42
        wb.save(p)

        # Must reject under a 1 KiB cap.
        tight = _limits(max_xlsx_uncompressed_bytes=1024)
        with catch_exceptions() as excs:
            safe_load_workbook(p, tight)
        assert excs and isinstance(excs[0], XlsxTooLargeError)

        # Must load cleanly under default limits.
        result = safe_load_workbook(p, SecurityLimits.defaults())
        assert result is not None
