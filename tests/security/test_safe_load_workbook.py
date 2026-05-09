"""Tests for safe_load_workbook: zip walk, streaming cell-count, ratio bomb."""

from __future__ import annotations

import dataclasses
import tempfile
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
from testsweet import catch_exceptions, test

from sheetwright.exceptions import XlsxTooLargeError
from sheetwright.security import SecurityLimits
from sheetwright.xlsx.safe_load import safe_load_workbook


def _limits(**kwargs: Any) -> SecurityLimits:
    return dataclasses.replace(SecurityLimits.defaults(), **kwargs)


def _write_simple_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws['A1'] = 'hello'
    ws['B2'] = 42
    wb.save(path)


def _write_xlsx_with_n_sheets(path: Path, n: int) -> None:
    wb = openpyxl.Workbook()
    wb.active.title = 'Sheet1'
    for i in range(2, n + 1):
        wb.create_sheet(f'Sheet{i}')
    wb.save(path)


def _write_xlsx_with_cells(path: Path, n_cells: int) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    for i in range(n_cells):
        ws.cell(row=i + 1, column=1, value=i)
    wb.save(path)


# ---------------------------------------------------------------------------
# Byte cap
# ---------------------------------------------------------------------------


@test
def byte_cap_rejects_oversized_xlsx():
    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td)
        p = tmp_path / 'big.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        for i in range(200):
            ws.cell(row=i + 1, column=1, value='x' * 100)
        wb.save(p)
        limits = _limits(max_xlsx_uncompressed_bytes=1024)
        with catch_exceptions() as excs:
            safe_load_workbook(p, limits)
        assert excs and isinstance(excs[0], XlsxTooLargeError)
        assert 'max_xlsx_uncompressed_bytes' in str(excs[0])


@test
def byte_cap_passes_under_default_cap():
    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td)
        p = tmp_path / 'small.xlsx'
        _write_simple_xlsx(p)
        wb = safe_load_workbook(p, SecurityLimits.defaults())
        assert wb is not None


# ---------------------------------------------------------------------------
# Sheet count cap
# ---------------------------------------------------------------------------


@test
def sheet_count_cap_rejects_too_many_sheets():
    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td)
        p = tmp_path / 'many.xlsx'
        _write_xlsx_with_n_sheets(p, 5)
        limits = _limits(max_xlsx_sheet_count=2)
        with catch_exceptions() as excs:
            safe_load_workbook(p, limits)
        assert excs and isinstance(excs[0], XlsxTooLargeError)
        assert 'max_xlsx_sheet_count' in str(excs[0])


@test
def sheet_count_cap_passes_under_cap():
    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td)
        p = tmp_path / 'few.xlsx'
        _write_xlsx_with_n_sheets(p, 5)
        limits = _limits(max_xlsx_sheet_count=10)
        wb = safe_load_workbook(p, limits)
        assert wb is not None


# ---------------------------------------------------------------------------
# Streaming cell-count (read_only=True path)
# ---------------------------------------------------------------------------


@test
def streaming_cell_count_rejects_over_cap():
    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td)
        p = tmp_path / 'cells.xlsx'
        _write_xlsx_with_cells(p, 100)
        limits = _limits(max_xlsx_cells_per_sheet=50)
        with catch_exceptions() as excs:
            safe_load_workbook(p, limits, read_only=True)
        assert excs and isinstance(excs[0], XlsxTooLargeError)
        assert 'max_xlsx_cells_per_sheet' in str(excs[0])


@test
def streaming_cell_count_passes_under_cap():
    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td)
        p = tmp_path / 'cells.xlsx'
        _write_xlsx_with_cells(p, 100)
        limits = _limits(max_xlsx_cells_per_sheet=200)
        wb = safe_load_workbook(p, limits, read_only=True)
        try:
            assert wb is not None
        finally:
            wb.close()


# ---------------------------------------------------------------------------
# Compression-ratio bomb boundary tests
#
# Use _FakeZipFile to fabricate ZipInfo entries with controlled
# file_size / compress_size. Real zip I/O is not needed for boundary tests.
# ---------------------------------------------------------------------------


class _FakeZipFile:
    """Context-manager wrapper that returns fabricated ZipInfo entries."""

    def __init__(self, path: Path, entries: list[zipfile.ZipInfo]):
        self._path = path
        self._entries = entries

    def __enter__(self) -> '_FakeZipFile':
        return self

    def __exit__(self, *_: object) -> None:
        pass

    def infolist(self) -> list[zipfile.ZipInfo]:
        return self._entries


def _fake_info(
    filename: str, file_size: int, compress_size: int
) -> zipfile.ZipInfo:
    info = zipfile.ZipInfo(filename)
    info.file_size = file_size
    info.compress_size = compress_size
    return info


def _run_check_zip_with_fake(
    entries: list[zipfile.ZipInfo], limits: SecurityLimits
) -> Exception | None:
    """Run _check_zip using a patched zipfile.ZipFile that returns ``entries``."""
    from unittest.mock import patch

    from sheetwright.xlsx import safe_load as safe_load_mod

    fake = _FakeZipFile(Path('fake.xlsx'), entries)

    with patch.object(zipfile, 'ZipFile', return_value=fake):
        try:
            safe_load_mod._check_zip(Path('fake.xlsx'), limits)
            return None
        except XlsxTooLargeError as e:
            return e


_8_MIB = 8 * 1024 * 1024
_16_MIB = 16 * 1024 * 1024
_4_MIB = 4 * 1024 * 1024


@test
def ratio_bomb_high_ratio_small_file_passes():
    """Ratio 1500, file_size 4 MiB → should pass (below size floor)."""
    info = _fake_info('xl/worksheets/sheet1.xml', _4_MIB, _4_MIB // 1500)
    exc = _run_check_zip_with_fake([info], SecurityLimits.defaults())
    assert exc is None, f'Expected no error but got: {exc}'


@test
def ratio_bomb_high_ratio_large_file_raises():
    """Ratio 1500, file_size 16 MiB → should raise (both conditions met)."""
    info = _fake_info('xl/worksheets/sheet1.xml', _16_MIB, _16_MIB // 1500)
    exc = _run_check_zip_with_fake([info], SecurityLimits.defaults())
    assert exc is not None and isinstance(exc, XlsxTooLargeError), (
        f'Expected XlsxTooLargeError but got: {exc}'
    )


@test
def ratio_bomb_low_ratio_large_file_passes():
    """Ratio 500, file_size 16 MiB → should pass (ratio below threshold)."""
    info = _fake_info('xl/worksheets/sheet1.xml', _16_MIB, _16_MIB // 500)
    exc = _run_check_zip_with_fake([info], SecurityLimits.defaults())
    assert exc is None, f'Expected no error but got: {exc}'


# ---------------------------------------------------------------------------
# Shared-strings cap
# ---------------------------------------------------------------------------


@test
def shared_strings_cap_rejects_large_shared_strings():
    """Use a very low max_xlsx_shared_strings so a small file triggers it."""
    limits = _limits(max_xlsx_shared_strings=1)
    # cap = 1 * 32 = 32 bytes; any real sharedStrings.xml will exceed this
    # Build a fake ZipInfo for xl/sharedStrings.xml with file_size=100
    info = _fake_info('xl/sharedStrings.xml', 100, 50)
    exc = _run_check_zip_with_fake([info], limits)
    assert exc is not None and isinstance(exc, XlsxTooLargeError), (
        f'Expected XlsxTooLargeError but got: {exc}'
    )
    assert 'sharedStrings' in str(exc)


@test
def shared_strings_cap_passes_under_cap():
    """Shared strings file within cap does not raise."""
    limits = _limits(max_xlsx_shared_strings=1_000_000)
    # cap = 32_000_000 bytes; 100 bytes well under
    info = _fake_info('xl/sharedStrings.xml', 100, 50)
    exc = _run_check_zip_with_fake([info], limits)
    assert exc is None, f'Expected no error but got: {exc}'


# ---------------------------------------------------------------------------
# Sanity: normal xlsx loads cleanly at defaults
# ---------------------------------------------------------------------------


@test
def sanity_small_xlsx_loads_at_defaults():
    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td)
        p = tmp_path / 'sanity.xlsx'
        _write_simple_xlsx(p)
        wb = safe_load_workbook(p, SecurityLimits.defaults())
        assert any(ws.title == 'Sheet1' for ws in wb.worksheets)
