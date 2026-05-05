import tempfile
from contextlib import contextmanager
from pathlib import Path

import openpyxl
from testsweet import test

from sheetwright.xlsx.reader import read_xlsx
from sheetwright.xlsx.writer import write_xlsx
from sheetwright.security import SecurityLimits


@contextmanager
def _tmp_path():
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


def _wb_with_freeze(path: Path, freeze: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    ws['A1'] = 'header'
    ws.freeze_panes = freeze
    wb.save(path)


@test
def xlsx_reader_picks_up_freeze_panes():
    with _tmp_path() as tmp_path:
        src = tmp_path / 'in.xlsx'
        _wb_with_freeze(src, 'B2')
        wb = read_xlsx(src, limits=SecurityLimits.defaults())
        assert wb.sheet('S').frozen_panes == 'B2'


@test
def xlsx_writer_emits_freeze_panes():
    with _tmp_path() as tmp_path:
        src = tmp_path / 'in.xlsx'
        _wb_with_freeze(src, 'C5')
        wb = read_xlsx(src, limits=SecurityLimits.defaults())
        out = tmp_path / 'out.xlsx'
        write_xlsx(wb, out)
        re_read = read_xlsx(out, limits=SecurityLimits.defaults())
        assert re_read.sheet('S').frozen_panes == 'C5'


@test
def no_freeze_round_trips_as_none():
    with _tmp_path() as tmp_path:
        src = tmp_path / 'in.xlsx'
        wb = openpyxl.Workbook()
        wb.active.title = 'S'
        wb.save(src)
        out = tmp_path / 'out.xlsx'
        write_xlsx(read_xlsx(src, limits=SecurityLimits.defaults()), out)
        assert (
            read_xlsx(out, limits=SecurityLimits.defaults())
            .sheet('S')
            .frozen_panes
            is None
        )
