"""Helpers that build small openpyxl workbooks for round-trip tests.

These are helpers callable from tests. We keep them in one place so test
setup stays consistent.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName


def write_simple_xlsx(path: Path) -> None:
    """A two-sheet workbook with values, a formula, and a named range."""
    wb = openpyxl.Workbook()
    s1 = wb.active
    s1.title = 'Inputs'
    s1['A1'] = 'growth_rate'
    s1['B1'] = 0.04
    s1['A2'] = 'base_revenue'
    s1['B2'] = 1_000_000

    s2 = wb.create_sheet('Outputs')
    s2['A1'] = 'revenue_2027'
    s2['B1'] = '=Inputs!B2 * (1 + Inputs!B1)'

    wb.defined_names['growth_rate'] = DefinedName(
        name='growth_rate', attr_text='Inputs!$B$1'
    )

    wb.save(path)


def write_formatted_xlsx(path: Path) -> None:
    """A workbook exercising fonts, fills, borders, number formats, column widths."""
    from openpyxl.styles import (
        Border as XBorder,
        Font as XFont,
        PatternFill,
        Side as XSide,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Format'

    ws['A1'] = 'bold red'
    ws['A1'].font = XFont(name='Calibri', size=11, bold=True, color='FFFF0000')

    ws['B1'] = 0.1234
    ws['B1'].number_format = '0.00%'

    ws['C1'] = 1234.5
    ws['C1'].number_format = '#,##0.00'

    ws['D1'] = 'filled'
    ws['D1'].fill = PatternFill(fill_type='solid', fgColor='FFFFFF00')

    ws['E1'] = 'boxed'
    side = XSide(style='thin', color='FF000000')
    ws['E1'].border = XBorder(left=side, right=side, top=side, bottom=side)

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 10

    wb.save(path)


def write_validation_xlsx(path: Path) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation as XDV

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'V'

    dv_list = XDV(type='list', formula1='"yes,no,maybe"', allow_blank=True)
    dv_list.add('A1:A10')
    ws.add_data_validation(dv_list)

    dv_range = XDV(type='whole', operator='between', formula1=1, formula2=100)
    dv_range.add('B1:B10')
    ws.add_data_validation(dv_range)

    wb.save(path)


def write_tier2_xlsx(path: Path) -> None:
    """One workbook exercising frozen panes, print area, comments,
    conditional formatting, and a ListObject table."""
    from openpyxl.comments import Comment as XComment
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.table import (
        Table as XTable,
        TableColumn as XTableColumn,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    ws['A1'], ws['B1'], ws['C1'] = 'Region', 'Q1', 'Q2'
    ws['A2'], ws['B2'], ws['C2'] = 'North', 100, 200
    ws['A3'], ws['B3'], ws['C3'] = 'South', 150, 250

    ws.freeze_panes = 'B2'
    ws.print_area = 'A1:C3'
    ws['B2'].comment = XComment('Q1 forecast', 'Alice')

    ws.conditional_formatting.add(
        'B2:C3',
        CellIsRule(
            operator='greaterThan',
            formula=['100'],
            fill=PatternFill(fill_type='solid', start_color='FF00FF00'),
        ),
    )
    ws.conditional_formatting.add(
        'B2:C3',
        ColorScaleRule(
            start_type='min',
            start_color='FFFF0000',
            end_type='max',
            end_color='FF00FF00',
        ),
    )

    cols = [
        XTableColumn(id=1, name='Region'),
        XTableColumn(id=2, name='Q1'),
        XTableColumn(id=3, name='Q2'),
    ]
    ws.add_table(
        XTable(
            displayName='Sales',
            name='Sales',
            ref='A1:C3',
            headerRowCount=1,
            totalsRowCount=0,
            tableColumns=cols,
        )
    )

    wb.save(path)
