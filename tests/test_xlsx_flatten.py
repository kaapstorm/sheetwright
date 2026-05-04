import tempfile
from contextlib import contextmanager
from pathlib import Path

from testsweet import test

from sheetwright.xlsx.flatten import (
    detect_external_refs,
    flatten_external_refs,
)
from sheetwright.xlsx.reader import read_xlsx
from tests.fixtures.external_xlsx import write_xlsx_with_external_ref


@contextmanager
def _tmp_path():
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


@test
def detect_finds_external_ref():
    with _tmp_path() as tmp_path:
        p = tmp_path / 'in.xlsx'
        write_xlsx_with_external_ref(p, cached_value=42.0)
        refs = detect_external_refs(p)
        assert any('[other.xlsx]' in r for r in refs)


@test
def detect_ignores_structured_table_references():
    """Plan 3 added ListObject tables. `=Sales[Region]` is a structured
    reference, NOT an external reference — must not be flagged."""
    import openpyxl
    from openpyxl.worksheet.table import Table, TableColumn

    with _tmp_path() as tmp_path:
        p = tmp_path / 'in.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'S'
        ws['A1'], ws['B1'] = 'Region', 'Q1'
        ws['A2'], ws['B2'] = 'North', 100
        ws.add_table(
            Table(
                displayName='Sales',
                name='Sales',
                ref='A1:B2',
                tableColumns=[
                    TableColumn(id=1, name='Region'),
                    TableColumn(id=2, name='Q1'),
                ],
            )
        )
        ws['C1'] = '=SUM(Sales[Q1])'
        wb.save(p)
        refs = detect_external_refs(p)
        assert refs == ()


@test
def detect_returns_empty_when_no_externals():
    import openpyxl

    with _tmp_path() as tmp_path:
        p = tmp_path / 'in.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = '=B1+1'
        wb.save(p)
        assert detect_external_refs(p) == ()


@test
def flatten_replaces_external_formula_with_cached_value():
    with _tmp_path() as tmp_path:
        src = tmp_path / 'in.xlsx'
        write_xlsx_with_external_ref(src, cached_value=42.0)
        wb = read_xlsx(src)
        flatten_external_refs(wb, src)
        cell = wb.sheet('S').get('A1')
        assert cell.formula is None
        assert cell.value == 42.0


@test
def flatten_leaves_internal_formulas_untouched():
    """Only external-ref formulas should flatten; internal formulas
    stay as formulas."""
    import openpyxl

    with _tmp_path() as tmp_path:
        src = tmp_path / 'in.xlsx'
        wb_op = openpyxl.Workbook()
        ws = wb_op.active
        ws.title = 'S'
        ws['A1'] = 1
        ws['A2'] = '=A1*2'
        wb_op.save(src)
        wb = read_xlsx(src)
        flatten_external_refs(wb, src)
        assert wb.sheet('S').get('A2').formula == '=A1*2'
