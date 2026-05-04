from pathlib import Path

import pytest
from unmagic import fixture, use

from claudesheets.mcp.errors import MCPError
from claudesheets.mcp.server import (
    do_reimport_abort,
    do_reimport_apply,
    do_reimport_stage,
)
from tests.fixtures.workbooks import write_simple_xlsx


@fixture
def populated(tmp_path: Path):
    src = tmp_path / 'in.xlsx'
    write_simple_xlsx(src)
    p = tmp_path / 'proj'
    p.mkdir()
    (p / 'claudesheets.toml').write_text(
        '[project]\nname = "in"\n[build]\ncalc_engine = "libreoffice"\n'
    )
    (p / 'workbook.toml').write_text('[workbook]\nname = "in"\nsheets = []\n')
    (p / 'sheets').mkdir()
    (p / 'data').mkdir()
    from click.testing import CliRunner

    from claudesheets.cli import main

    CliRunner().invoke(main, ['import', str(src), '--project', str(p)])
    yield p, src


@use(populated)
def test_reimport_stage_returns_empty_for_unchanged_xlsx(tmp_path: Path):
    p, src = populated()
    out = do_reimport_stage(
        xlsx=str(src), project=str(p), flatten=False, force=False
    )
    assert out['is_empty'] is True
    assert not (p / '.claudesheets' / 'reimport.json').is_file()


@use(populated)
def test_reimport_stage_then_apply(tmp_path: Path):
    p, _ = populated()
    new_src = tmp_path / 'changed.xlsx'
    import openpyxl

    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'Inputs'
    s['A1'], s['B1'] = 'growth_rate', 0.99
    s['A2'], s['B2'] = 'base_revenue', 1_000_000
    wb.create_sheet('Outputs')
    wb.save(new_src)

    staged = do_reimport_stage(
        xlsx=str(new_src), project=str(p), flatten=False, force=False
    )
    assert staged['is_empty'] is False
    assert (p / '.claudesheets' / 'reimport.json').is_file()

    out = do_reimport_apply(project=str(p), archive=False)
    assert out['ok'] is True
    md = p / 'sheets' / '01_inputs.md'
    assert '0.99' in md.read_text()
    assert not (p / '.claudesheets' / 'reimport.json').is_file()


@use(populated)
def test_reimport_abort_clears_session(tmp_path: Path):
    p, _ = populated()
    new_src = tmp_path / 'changed.xlsx'
    import openpyxl

    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'Inputs'
    s['A1'], s['B1'] = 'growth_rate', 0.50
    s['A2'], s['B2'] = 'base_revenue', 1_000_000
    wb.create_sheet('Outputs')
    wb.save(new_src)

    do_reimport_stage(
        xlsx=str(new_src), project=str(p), flatten=False, force=False
    )
    out = do_reimport_abort(project=str(p))
    assert out['ok'] is True
    assert not (p / '.claudesheets' / 'reimport.json').is_file()


@use(populated)
def test_reimport_stage_then_restage_with_no_changes_clears_session(
    tmp_path: Path,
):
    """C2 regression: a no-changes restage must clear any prior session."""
    p, src = populated()

    # Stage with a modified xlsx — non-empty diff, session saved.
    new_src = tmp_path / 'changed.xlsx'
    import openpyxl

    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'Inputs'
    s['A1'], s['B1'] = 'growth_rate', 0.99
    s['A2'], s['B2'] = 'base_revenue', 1_000_000
    s2 = wb.create_sheet('Outputs')
    s2['A1'] = 'revenue_2027'
    s2['B1'] = '=Inputs!B2 * (1 + Inputs!B1)'
    wb.save(new_src)

    do_reimport_stage(
        xlsx=str(new_src), project=str(p), flatten=False, force=False
    )
    assert (p / '.claudesheets' / 'reimport.json').is_file()

    # Re-stage with the original xlsx — diff is empty; session must be cleared.
    out = do_reimport_stage(
        xlsx=str(src), project=str(p), flatten=False, force=False
    )
    assert out['is_empty'] is True
    assert not (p / '.claudesheets' / 'reimport.json').is_file()


@use(populated)
def test_reimport_apply_without_session_raises_typed_error(tmp_path: Path):
    p, _ = populated()
    with pytest.raises(MCPError) as excinfo:
        do_reimport_apply(project=str(p), archive=False)
    assert excinfo.value.code == 'no_staged_session'
