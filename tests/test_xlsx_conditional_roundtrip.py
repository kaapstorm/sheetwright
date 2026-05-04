import tempfile
from contextlib import contextmanager
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import (
    CellIsRule,
    ColorScaleRule,
    FormulaRule,
)
from openpyxl.styles import PatternFill
from testsweet import test

from claudesheets.xlsx.reader import read_xlsx
from claudesheets.xlsx.writer import write_xlsx


@contextmanager
def _tmp_path():
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


def _wb_with_cell_is(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    rule = CellIsRule(
        operator='greaterThan',
        formula=['0'],
        fill=PatternFill(fill_type='solid', start_color='FF00FF00'),
    )
    ws.conditional_formatting.add('A1:A10', rule)
    wb.save(path)


def _wb_with_formula(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    rule = FormulaRule(
        formula=['ISERROR(A1)'],
        fill=PatternFill(fill_type='solid', start_color='FFFF0000'),
    )
    ws.conditional_formatting.add('A1:A10', rule)
    wb.save(path)


def _wb_with_color_scale(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    rule = ColorScaleRule(
        start_type='min',
        start_color='FFFF0000',
        mid_type='percentile',
        mid_value=50,
        mid_color='FFFFFF00',
        end_type='max',
        end_color='FF00FF00',
    )
    ws.conditional_formatting.add('A1:A10', rule)
    wb.save(path)


@test
def cell_is_rule_round_trips():
    from claudesheets.model.conditional import CellIsRule

    with _tmp_path() as tmp_path:
        src = tmp_path / 'in.xlsx'
        _wb_with_cell_is(src)
        out = tmp_path / 'out.xlsx'
        write_xlsx(read_xlsx(src), out)
        s = read_xlsx(out).sheet('S')
        assert len(s.conditional_formats) == 1
        cf = s.conditional_formats[0]
        assert isinstance(cf, CellIsRule)
        assert 'A1:A10' in cf.ranges
        assert cf.operator == 'greaterThan'
        assert cf.formula == ('0',)


@test
def formula_rule_round_trips():
    from claudesheets.model.conditional import FormulaRule

    with _tmp_path() as tmp_path:
        src = tmp_path / 'in.xlsx'
        _wb_with_formula(src)
        out = tmp_path / 'out.xlsx'
        write_xlsx(read_xlsx(src), out)
        s = read_xlsx(out).sheet('S')
        assert any(isinstance(cf, FormulaRule) for cf in s.conditional_formats)
        cf = next(
            c for c in s.conditional_formats if isinstance(c, FormulaRule)
        )
        assert cf.formula == 'ISERROR(A1)'


@test
def color_scale_rule_round_trips():
    from claudesheets.model.conditional import ColorScaleRule

    with _tmp_path() as tmp_path:
        src = tmp_path / 'in.xlsx'
        _wb_with_color_scale(src)
        out = tmp_path / 'out.xlsx'
        write_xlsx(read_xlsx(src), out)
        s = read_xlsx(out).sheet('S')
        cf = next(
            c for c in s.conditional_formats if isinstance(c, ColorScaleRule)
        )
        assert cf.start_type == 'min'
        assert cf.end_type == 'max'
        assert cf.mid_type == 'percentile'
