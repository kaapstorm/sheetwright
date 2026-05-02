"""Translate openpyxl conditional-formatting rules <-> typed CF model.

Each `rule.type` (string from xlsx) maps to one of our five frozen
dataclasses; each dataclass maps back to the corresponding openpyxl
`*Rule` constructor. Unknown rule types raise — we don't silently
drop them.
"""

from __future__ import annotations

from typing import Optional, Tuple

from openpyxl.formatting.rule import (
    CellIsRule as XCellIsRule,
    ColorScaleRule as XColorScaleRule,
    DataBarRule as XDataBarRule,
    FormulaRule as XFormulaRule,
    IconSetRule as XIconSetRule,
    Rule,
)
from openpyxl.styles import PatternFill
from openpyxl.styles.differential import DifferentialStyle

from claudesheets.model.conditional import (
    CellIsRule,
    CFStyle,
    ColorScaleRule,
    ConditionalFormat,
    DataBarRule,
    FormulaRule,
    IconSetRule,
)


def cf_from_openpyxl_rule(
    ranges: Tuple[str, ...], rule: Rule
) -> ConditionalFormat:
    style = _style_from_dxf(rule.dxf) if rule.dxf else None
    common = {
        'ranges': ranges,
        'priority': rule.priority,
        'stop_if_true': bool(getattr(rule, 'stopIfTrue', False)),
    }

    if rule.type == 'cellIs':
        return CellIsRule(
            **common,  # type: ignore[arg-type]
            operator=rule.operator or 'equal',
            formula=tuple(rule.formula or ()),
            style=style,
        )
    if rule.type == 'expression':
        formulas = list(rule.formula or [])
        return FormulaRule(
            **common,  # type: ignore[arg-type]
            formula=formulas[0] if formulas else '',
            style=style,
        )
    if rule.type == 'colorScale' and rule.colorScale is not None:
        cs = rule.colorScale
        cfvo = list(cs.cfvo or [])
        colors = [_color_value(c) for c in (cs.color or [])]

        def _at(seq, i, default=None):
            return seq[i] if i < len(seq) else default

        if len(cfvo) == 3:
            return ColorScaleRule(
                **common,  # type: ignore[arg-type]
                start_type=cfvo[0].type,
                start_value=_str_or_none(cfvo[0].val),
                start_color=_at(colors, 0, 'FFFFFFFF'),
                mid_type=cfvo[1].type,
                mid_value=_str_or_none(cfvo[1].val),
                mid_color=_at(colors, 1),
                end_type=cfvo[2].type,
                end_value=_str_or_none(cfvo[2].val),
                end_color=_at(colors, 2, 'FF000000'),
            )
        return ColorScaleRule(
            **common,  # type: ignore[arg-type]
            start_type=cfvo[0].type if cfvo else 'min',
            start_value=_str_or_none(cfvo[0].val) if cfvo else None,
            start_color=_at(colors, 0, 'FFFFFFFF'),
            end_type=cfvo[1].type if len(cfvo) > 1 else 'max',
            end_value=(_str_or_none(cfvo[1].val) if len(cfvo) > 1 else None),
            end_color=_at(colors, 1, 'FF000000'),
        )
    if rule.type == 'dataBar' and rule.dataBar is not None:
        db = rule.dataBar
        cfvo = list(db.cfvo or [])
        return DataBarRule(
            **common,  # type: ignore[arg-type]
            start_type=cfvo[0].type if cfvo else 'min',
            start_value=_str_or_none(cfvo[0].val) if cfvo else None,
            end_type=cfvo[1].type if len(cfvo) > 1 else 'max',
            end_value=(_str_or_none(cfvo[1].val) if len(cfvo) > 1 else None),
            color=_color_value(db.color) if db.color else 'FF638EC6',
            show_value=(
                bool(db.showValue) if db.showValue is not None else True
            ),
        )
    if rule.type == 'iconSet' and rule.iconSet is not None:
        ic = rule.iconSet
        return IconSetRule(
            **common,  # type: ignore[arg-type]
            icon_style=ic.iconSet or '3TrafficLights1',
            type=ic.cfvo[0].type if ic.cfvo else 'percent',
            values=tuple(_str_or_none(c.val) or '0' for c in (ic.cfvo or [])),
        )

    raise ValueError(
        f'unsupported conditional format rule type: {rule.type!r}'
    )


def cf_to_openpyxl_rule(cf: ConditionalFormat) -> Rule:
    if isinstance(cf, CellIsRule):
        return XCellIsRule(
            operator=cf.operator,
            formula=list(cf.formula),
            fill=_fill_from_style(cf.style),
            stopIfTrue=cf.stop_if_true,
        )
    if isinstance(cf, FormulaRule):
        return XFormulaRule(
            formula=[cf.formula] if cf.formula else [],
            fill=_fill_from_style(cf.style),
            stopIfTrue=cf.stop_if_true,
        )
    if isinstance(cf, ColorScaleRule):
        return XColorScaleRule(
            start_type=cf.start_type,
            start_value=cf.start_value,
            start_color=cf.start_color,
            mid_type=cf.mid_type,
            mid_value=cf.mid_value,
            mid_color=cf.mid_color,
            end_type=cf.end_type,
            end_value=cf.end_value,
            end_color=cf.end_color,
        )
    if isinstance(cf, DataBarRule):
        return XDataBarRule(
            start_type=cf.start_type,
            start_value=cf.start_value,
            end_type=cf.end_type,
            end_value=cf.end_value,
            color=cf.color,
            showValue=cf.show_value,
        )
    if isinstance(cf, IconSetRule):
        return XIconSetRule(
            icon_style=cf.icon_style,
            type=cf.type,
            values=list(cf.values),
        )
    raise TypeError(f'unknown ConditionalFormat type: {type(cf).__name__}')


def _style_from_dxf(dxf: DifferentialStyle) -> Optional[CFStyle]:
    fill_color: Optional[str] = None
    if dxf.fill is not None:
        fg = getattr(dxf.fill, 'fgColor', None)
        if fg is not None:
            fill_color = fg.value or fg.rgb or None

    bold = italic = False
    font_color: Optional[str] = None
    if dxf.font is not None:
        bold = bool(dxf.font.b)
        italic = bool(dxf.font.i)
        if dxf.font.color is not None:
            fc = dxf.font.color.value or dxf.font.color.rgb or None
            font_color = str(fc) if fc is not None else None

    if not (fill_color or bold or italic or font_color):
        return None
    return CFStyle(
        fill_color=fill_color,
        font_bold=bold,
        font_italic=italic,
        font_color=font_color,
    )


def _fill_from_style(style: Optional[CFStyle]) -> Optional[PatternFill]:
    if style is None or not style.fill_color:
        return None
    return PatternFill(fill_type='solid', start_color=style.fill_color)


def _color_value(c: object) -> str:
    """Best-effort extraction of an openpyxl color to a hex string."""
    return getattr(c, 'value', None) or getattr(c, 'rgb', None) or ''


def _str_or_none(v: object) -> Optional[str]:
    if v is None:
        return None
    return str(v)
