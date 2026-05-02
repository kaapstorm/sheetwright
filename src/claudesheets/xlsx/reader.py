"""Read an .xlsx file into a Workbook."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any, Optional

import openpyxl
from openpyxl.cell.cell import Cell as XCell
from openpyxl.worksheet.datavalidation import DataValidation as XDV

from claudesheets.model.cell import Cell
from claudesheets.model.comment import Comment as Cmt
from claudesheets.model.conditional import ConditionalFormat
from claudesheets.model.format import Border, CellFormat, Fill, Font, Side
from claudesheets.model.table import ListTable, ListTableColumn
from claudesheets.model.validation import DataValidation
from claudesheets.model.workbook import NamedRange, Sheet, Workbook
from claudesheets.xlsx.cf_translate import cf_from_openpyxl_rule


_PRINT_AREA_PREFIX = re.compile(r"^(?:'[^']+'|[^!]+)!")


def _strip_sheet_prefix_and_dollars(area: Optional[str]) -> Optional[str]:
    if not area:
        return None
    s = _PRINT_AREA_PREFIX.sub('', area)
    return s.replace('$', '')


def _normalise_color(c: object) -> Optional[str]:
    """openpyxl colours can be Color objects, theme refs, or hex strings.

    For Plan 1 we only carry hex RGB. Theme colours and indexed colours are
    dropped (None). Alpha-channel hex (8 chars) is normalised to 6-char RGB.
    """
    if c is None:
        return None
    rgb = getattr(c, 'rgb', None) or (c if isinstance(c, str) else None)
    if not isinstance(rgb, str):
        return None
    if len(rgb) == 8:  # AARRGGBB
        return rgb[2:].upper()
    if len(rgb) == 6:
        return rgb.upper()
    return None


def _read_font(c: XCell) -> Optional[Font]:
    f = c.font
    if not f:
        return None
    has_anything = any(
        [
            f.name,
            f.size,
            f.bold,
            f.italic,
            f.underline,
            _normalise_color(f.color),
        ]
    )
    if not has_anything:
        return None
    return Font(
        name=f.name or None,
        size=float(f.size) if f.size is not None else None,
        bold=bool(f.bold),
        italic=bool(f.italic),
        underline=f.underline if f.underline in ('single', 'double') else None,
        color=_normalise_color(f.color),
    )


def _read_fill(c: XCell) -> Optional[Fill]:
    fill = c.fill
    if not fill or fill.fill_type != 'solid':
        return None
    color = _normalise_color(fill.fgColor)
    if not color:
        return None
    return Fill(color=color)


def _read_side(s: Any) -> Optional[Side]:
    if not s or not s.style:
        return None
    return Side(style=s.style, color=_normalise_color(s.color))


def _read_border(c: XCell) -> Optional[Border]:
    b = c.border
    if not b:
        return None
    sides = {
        n: _read_side(getattr(b, n))
        for n in ('left', 'right', 'top', 'bottom')
    }
    if not any(sides.values()):
        return None
    return Border(**sides)


def _cell_format(c: XCell) -> Optional[CellFormat]:
    font = _read_font(c)
    fill = _read_fill(c)
    border = _read_border(c)
    nf = (
        c.number_format
        if c.number_format and c.number_format != 'General'
        else None
    )
    if not any([font, fill, border, nf]):
        return None
    return CellFormat(font=font, fill=fill, border=border, number_format=nf)


def _read_validation(dv: XDV) -> DataValidation:
    ranges = [str(r) for r in dv.sqref.ranges] if dv.sqref else []
    return DataValidation(
        type=str(dv.type) if dv.type else '',
        ranges=ranges,
        operator=str(dv.operator) if dv.operator else None,
        formula1=dv.formula1,
        formula2=dv.formula2,
        allow_blank=bool(dv.allow_blank),
    )


def _format_id(fmt: CellFormat) -> str:
    """Stable, content-addressed id for a CellFormat."""
    h = hashlib.sha1(repr(fmt).encode('utf-8')).hexdigest()[:10]
    return f'f-{h}'


def _read_tables(ws: Any) -> list[ListTable]:
    out: list[ListTable] = []
    for tbl in ws.tables.values():
        cols = tuple(
            ListTableColumn(
                name=tc.name,
                formula=tc.calculatedColumnFormula,
                totals_label=tc.totalsRowLabel,
                totals_function=tc.totalsRowFunction,
            )
            for tc in (tbl.tableColumns or [])
        )
        out.append(
            ListTable(
                name=tbl.displayName,
                ref=tbl.ref,
                header_row_count=(
                    tbl.headerRowCount if tbl.headerRowCount is not None else 1
                ),
                totals_row_count=tbl.totalsRowCount or 0,
                columns=cols,
                style=(
                    tbl.tableStyleInfo.name if tbl.tableStyleInfo else None
                ),
            )
        )
    return out


def _read_conditional_formats(ws: Any) -> list[ConditionalFormat]:
    out: list[ConditionalFormat] = []
    for item in ws.conditional_formatting:
        ranges = tuple(str(r) for r in item.sqref.ranges)
        for rule in item.rules:
            out.append(cf_from_openpyxl_rule(ranges, rule))
    return out


def read_xlsx(path: Path) -> Workbook:
    path = Path(path)
    src = openpyxl.load_workbook(path, data_only=False)
    wb = Workbook(name=path.stem)

    for ws in src.worksheets:
        sheet = Sheet(name=ws.title)
        for col_letter, dim in ws.column_dimensions.items():
            if dim.width is not None:
                sheet.column_widths[col_letter] = float(dim.width)
        for dv in ws.data_validations.dataValidation:
            sheet.validations.append(_read_validation(dv))
        sheet.frozen_panes = ws.freeze_panes
        sheet.print_area = _strip_sheet_prefix_and_dollars(ws.print_area)
        sheet.tables.extend(_read_tables(ws))
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c, XCell):
                    # MergedCell — skip; no independent value or format
                    continue
                fmt = _cell_format(c)
                fmt_id = None
                if fmt is not None:
                    fmt_id = _format_id(fmt)
                    sheet.formats[fmt_id] = fmt

                if c.value is None and fmt_id is None:
                    continue
                if isinstance(c.value, str) and c.value.startswith('='):
                    sheet.set(
                        c.coordinate, Cell(formula=c.value, format_id=fmt_id)
                    )
                elif c.value is None:
                    # blank but formatted
                    sheet.set(c.coordinate, Cell(format_id=fmt_id))
                else:
                    # openpyxl stub types c.value more widely than CellValue
                    # (Decimal, RichText, etc.) — these will be narrowed in
                    # later tasks; for now pass through.
                    sheet.set(
                        c.coordinate,
                        Cell(value=c.value, format_id=fmt_id),  # type: ignore[arg-type]
                    )
                if c.comment is not None:
                    sheet.comments[c.coordinate] = Cmt(
                        author=c.comment.author or '',
                        text=c.comment.text or '',
                    )

        sheet.conditional_formats.extend(_read_conditional_formats(ws))
        wb.sheets.append(sheet)

    for name, defn in src.defined_names.items():
        wb.named_ranges.append(
            NamedRange(name=name, ref=defn.attr_text, scope='workbook')
        )
    for ws in src.worksheets:
        for name, defn in ws.defined_names.items():
            wb.named_ranges.append(
                NamedRange(
                    name=name,
                    ref=defn.attr_text,
                    scope='sheet',
                    sheet=ws.title,
                )
            )

    return wb
