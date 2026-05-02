"""Write a Workbook to an .xlsx file."""

from __future__ import annotations

import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Literal, Optional

import openpyxl
from openpyxl.cell.cell import Cell as XCell
from openpyxl.comments import Comment as XComment
from openpyxl.worksheet.datavalidation import DataValidation as XDV
from openpyxl.worksheet.table import (
    Table as XTable,
    TableColumn as XTableColumn,
    TableStyleInfo as XTableStyleInfo,
)
from openpyxl.styles import (
    Border as XBorder,
    Color,
    Font as XFont,
    PatternFill,
    Side as XSide,
)
from openpyxl.workbook.defined_name import DefinedName

from claudesheets.model.format import Border, CellFormat, Fill, Font, Side
from claudesheets.model.workbook import Workbook

# openpyxl's underline Literal type
_Underline = Optional[
    Literal['single', 'double', 'singleAccounting', 'doubleAccounting', 'none']
]

# Deterministic epoch for all generated xlsx files
_DETERMINISTIC_EPOCH = datetime(2000, 1, 1)


def _fix_xlsx_timestamps(path: Path, epoch: datetime) -> None:
    """Rewrite xlsx core.xml and zip metadata for deterministic builds.

    openpyxl sets properties.modified to the current time during save
    regardless of what we've set. Also, zip file entry metadata includes
    timestamps. We post-process both to ensure determinism.
    """
    # Convert epoch to DOS timestamp (year, month, day, hour, minute, second)
    dos_date_time = (
        epoch.year,
        epoch.month,
        epoch.day,
        epoch.hour,
        epoch.minute,
        epoch.second,
    )

    # Read all files from the zip, preserving entry order
    with zipfile.ZipFile(path, 'r') as z:
        order = z.namelist()
        files_data = {name: z.read(name) for name in order}

    # Fix core.xml timestamp
    core_xml = files_data['docProps/core.xml'].decode('utf-8')
    epoch_str = epoch.isoformat() + 'Z'
    core_xml, n = re.subn(
        r'<dcterms:modified[^>]*>.*?</dcterms:modified>',
        (
            f'<dcterms:modified xsi:type="dcterms:W3CDTF">'
            f'{epoch_str}</dcterms:modified>'
        ),
        core_xml,
    )
    if n != 1:
        raise RuntimeError(
            f'expected exactly one <dcterms:modified> in core.xml, '
            f'found {n}; openpyxl output may have changed'
        )
    files_data['docProps/core.xml'] = core_xml.encode('utf-8')

    # Rewrite the zip with deterministic timestamps, preserving order
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for name in order:
            info = zipfile.ZipInfo(name, dos_date_time)
            # Use DEFLATE compression for all files (except dirs)
            if not name.endswith('/'):
                info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, files_data[name])


def _xfont(font: Optional[Font]) -> Optional[XFont]:
    if font is None:
        return None
    color = Color(rgb=('FF' + font.color)) if font.color else None
    underline: _Underline = font.underline  # type: ignore[assignment]
    return XFont(
        name=font.name,
        size=font.size,
        bold=font.bold,
        italic=font.italic,
        underline=underline,
        color=color,
    )


def _xfill(fill: Optional[Fill]) -> Optional[PatternFill]:
    if fill is None or not fill.color:
        return None
    return PatternFill(
        fill_type='solid', fgColor=Color(rgb=('FF' + fill.color))
    )


def _xside(side: Optional[Side]) -> Optional[XSide]:
    if side is None or not side.style:
        return None
    color = Color(rgb=('FF' + side.color)) if side.color else None
    return XSide(style=side.style, color=color)  # type: ignore[arg-type]


def _xborder(border: Optional[Border]) -> Optional[XBorder]:
    if border is None:
        return None
    return XBorder(
        left=_xside(border.left),
        right=_xside(border.right),
        top=_xside(border.top),
        bottom=_xside(border.bottom),
    )


def _apply_format(cell: XCell, fmt: CellFormat) -> None:
    if fmt.font is not None:
        cell.font = _xfont(fmt.font)  # type: ignore[assignment]
    if fmt.fill is not None:
        cell.fill = _xfill(fmt.fill)  # type: ignore[assignment]
    if fmt.border is not None:
        cell.border = _xborder(fmt.border)  # type: ignore[assignment]
    if fmt.number_format is not None:
        cell.number_format = fmt.number_format


def write_xlsx(wb: Workbook, path: Path) -> None:
    path = Path(path)
    out = openpyxl.Workbook()
    default = out.active
    if default is not None:
        out.remove(default)

    for sheet in wb.sheets:
        ws = out.create_sheet(title=sheet.name)

        for col, width in sheet.column_widths.items():
            ws.column_dimensions[col].width = width

        for v in sheet.validations:
            xdv = XDV(
                type=v.type,  # type: ignore[arg-type]
                operator=v.operator,  # type: ignore[arg-type]
                formula1=v.formula1,
                formula2=v.formula2,
                allow_blank=v.allow_blank,
            )
            for r in v.ranges:
                xdv.add(r)
            ws.add_data_validation(xdv)

        if sheet.frozen_panes:
            ws.freeze_panes = sheet.frozen_panes

        if sheet.print_area:
            ws.print_area = sheet.print_area

        for t in sheet.tables:
            xcols = [
                XTableColumn(
                    id=i + 1,
                    name=c.name,
                    calculatedColumnFormula=c.formula,  # type: ignore[arg-type]
                    totalsRowLabel=c.totals_label,
                    totalsRowFunction=c.totals_function,  # type: ignore[arg-type]
                )
                for i, c in enumerate(t.columns)
            ]
            xt = XTable(
                displayName=t.name,
                name=t.name,
                ref=t.ref,
                headerRowCount=t.header_row_count,
                totalsRowCount=t.totals_row_count,
                tableColumns=xcols,
            )
            if t.style:
                xt.tableStyleInfo = XTableStyleInfo(name=t.style)
            ws.add_table(xt)

        from claudesheets.xlsx.cf_translate import cf_to_openpyxl_rule

        for cf in sheet.conditional_formats:
            xrule = cf_to_openpyxl_rule(cf)
            for r in cf.ranges:
                ws.conditional_formatting.add(r, xrule)

        for addr, cell in sheet.cells.items():
            xc = ws[addr]
            if cell.formula is not None:
                xc.value = cell.formula
            elif cell.value is not None:
                xc.value = cell.value
            if cell.format_id and cell.format_id in sheet.formats:
                _apply_format(xc, sheet.formats[cell.format_id])
            if addr in sheet.comments:
                cmt = sheet.comments[addr]
                ws[addr].comment = XComment(cmt.text, cmt.author)

    for nr in wb.named_ranges:
        defn = DefinedName(name=nr.name, attr_text=nr.ref)
        if nr.scope == 'workbook':
            out.defined_names[nr.name] = defn
        else:
            assert nr.sheet is not None
            out[nr.sheet].defined_names[nr.name] = defn

    out.properties.created = _DETERMINISTIC_EPOCH
    out.properties.modified = _DETERMINISTIC_EPOCH
    out.save(path)
    _fix_xlsx_timestamps(path, _DETERMINISTIC_EPOCH)
