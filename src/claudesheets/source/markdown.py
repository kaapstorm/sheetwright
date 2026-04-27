"""Parse and serialise a Sheet's value cells as a Markdown table.

Format:
- First row: `| (cell) | A | B | C | ...` where the second-onwards columns
  are the column letters. The first column header is the literal string
  `(cell)` and is reserved for the row index.
- Second row: GitHub-style separator `| --- | --- | --- | --- |`.
- Subsequent rows: `| <row-number> | <cell-A-content> | <cell-B-content> | ...`.

Cell content rules:
- Empty cell -> empty string.
- Strings are written verbatim, with `|` escaped as `\\|` and `\\` as `\\\\`.
- Numbers are written with repr() to preserve precision.
- Booleans are written as TRUE / FALSE (Excel's convention).
- Formula cells: the formula text, starting with `=`.

This module deliberately does NOT handle formats, validation, or
formula-keyed-by-name semantics — those live in the YAML sidecar.
"""

from __future__ import annotations

import re
from typing import List

from openpyxl.utils import column_index_from_string, get_column_letter

from claudesheets.model.cell import Cell
from claudesheets.model.workbook import Sheet

_HEADER = '(cell)'


def _escape(s: str) -> str:
    return s.replace('\\', '\\\\').replace('|', '\\|')


def _unescape(s: str) -> str:
    out: List[str] = []
    i = 0
    while i < len(s):
        c = s[i]
        if c == '\\' and i + 1 < len(s):
            out.append(s[i + 1])
            i += 2
            continue
        out.append(c)
        i += 1
    return ''.join(out)


def _format_value(cell: Cell) -> str:
    if cell.formula is not None:
        return _escape(cell.formula)
    v = cell.value
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, (int, float)):
        return repr(v)
    return _escape(str(v))


def _parse_value(raw: str) -> Cell:
    s = _unescape(raw.strip())
    if s == '':
        return Cell()
    if s.startswith('='):
        return Cell(formula=s)
    if s == 'TRUE':
        return Cell(value=True)
    if s == 'FALSE':
        return Cell(value=False)
    # Try int, then float, else string.
    if re.fullmatch(r'[+-]?\d+', s):
        try:
            return Cell(value=int(s))
        except ValueError:
            pass
    try:
        return Cell(value=float(s))
    except ValueError:
        return Cell(value=s)


def _parse_row(line: str) -> List[str]:
    # Split on unescaped `|`. Trim leading/trailing pipe.
    parts: List[str] = []
    buf: List[str] = []
    i = 0
    while i < len(line):
        c = line[i]
        if c == '\\' and i + 1 < len(line):
            buf.append(c)
            buf.append(line[i + 1])
            i += 2
            continue
        if c == '|':
            parts.append(''.join(buf))
            buf = []
            i += 1
            continue
        buf.append(c)
        i += 1
    parts.append(''.join(buf))
    # Trim the leading/trailing empty parts created by `|` at line edges.
    if parts and parts[0].strip() == '':
        parts = parts[1:]
    if parts and parts[-1].strip() == '':
        parts = parts[:-1]
    return parts


def dump_table(sheet: Sheet) -> str:
    if not sheet.cells:
        return f'| {_HEADER} |\n| --- |\n'

    max_row = 0
    max_col = 0
    for addr in sheet.cells:
        m = re.fullmatch(r'([A-Z]+)(\d+)', addr)
        if not m:
            raise ValueError(f'bad address: {addr!r}')
        col_letters, row_str = m.group(1), m.group(2)
        max_row = max(max_row, int(row_str))
        max_col = max(max_col, column_index_from_string(col_letters))

    headers = [_HEADER] + [get_column_letter(c) for c in range(1, max_col + 1)]
    sep = ['---'] * len(headers)

    lines = [
        '| ' + ' | '.join(headers) + ' |',
        '| ' + ' | '.join(sep) + ' |',
    ]
    for r in range(1, max_row + 1):
        row_cells = [str(r)]
        for c in range(1, max_col + 1):
            addr = f'{get_column_letter(c)}{r}'
            row_cells.append(_format_value(sheet.get(addr)))
        lines.append('| ' + ' | '.join(row_cells) + ' |')
    return '\n'.join(lines) + '\n'


def load_table(sheet: Sheet, text: str) -> None:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) < 2:
        return  # header only or empty

    header = _parse_row(lines[0])
    if not header or header[0].strip() != _HEADER:
        raise ValueError('first column header must be (cell)')
    column_letters = [h.strip() for h in header[1:]]

    for line in lines[2:]:  # skip header + separator
        cells = _parse_row(line)
        if not cells:
            continue
        row_idx_raw = cells[0].strip()
        if not row_idx_raw or not row_idx_raw.isdigit():
            continue
        row_idx = int(row_idx_raw)
        for i, raw in enumerate(cells[1:]):
            if i >= len(column_letters):
                break
            cell = _parse_value(raw)
            if cell.is_blank:
                continue
            sheet.set(f'{column_letters[i]}{row_idx}', cell)
