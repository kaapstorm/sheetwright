"""Structural lint over a Workbook + Project.

Detects:
  - dangling_sheet_ref: formula references a sheet that doesn't exist
  - dangling_named_range: formula uses an undefined name
  - manifest_sheet_missing: workbook.toml lists a sheet not on disk
  - sheet_file_missing_from_manifest: sheet file on disk not listed
  - bulk_data_missing: sheet's `source:` points at a missing csv

Formula parsing uses openpyxl's `Tokenizer`, which yields a structured
token stream (OPERAND/RANGE, FUNC, OP_*, etc.). This handles string
literals, nested function calls, and structured table references
(`Sales[Region]`) correctly -- none of which a regex-based parser
gets right.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import List, Optional, Tuple

from openpyxl.formula.tokenizer import Tokenizer

from sheetwright.config import load_workbook
from sheetwright.model.workbook import Workbook
from sheetwright.project import Project, slugify


@dataclass(frozen=True)
class CheckIssue:
    kind: str
    detail: str
    location: Optional[str] = None


def check_workbook(wb: Workbook, project: Project) -> List[CheckIssue]:
    issues: List[CheckIssue] = []
    sheet_names = {s.name for s in wb.sheets}
    table_names = {t.name for s in wb.sheets for t in s.tables}
    named = {nr.name for nr in wb.named_ranges}

    for sheet in wb.sheets:
        for addr, cell in sheet.cells.items():
            if cell.formula is None:
                continue
            sheet_refs, name_refs = _parse_formula_refs(cell.formula)
            for ref in sheet_refs:
                if ref not in sheet_names:
                    issues.append(
                        CheckIssue(
                            kind='dangling_sheet_ref',
                            detail=f'unknown sheet {ref!r}',
                            location=f'{sheet.name}!{addr}',
                        )
                    )
            for nm in name_refs:
                if nm in named or nm in table_names:
                    continue
                issues.append(
                    CheckIssue(
                        kind='dangling_named_range',
                        detail=f'unknown name {nm!r}',
                        location=f'{sheet.name}!{addr}',
                    )
                )

    # workbook.toml stores display names (e.g. "Inputs"); the
    # source reader/writer reconstruct stems via slugify+index. We
    # compare on the same stem keyspace.
    if project.workbook_toml.is_file():
        manifest = load_workbook(project.workbook_toml.read_text())
        manifest_stems = {
            f'{i:02d}_{slugify(name)}': name
            for i, name in enumerate(manifest.sheets, start=1)
        }
    else:
        manifest_stems = {}

    on_disk = {p.stem for p in project.sheets_dir.glob('*.md')}

    for stem in set(manifest_stems) - on_disk:
        name = manifest_stems[stem]
        issues.append(
            CheckIssue(
                kind='manifest_sheet_missing',
                detail=(
                    f'workbook.toml lists {name!r} but '
                    f'sheets/{stem}.md is missing'
                ),
                location='workbook.toml',
            )
        )
    for stem in on_disk - set(manifest_stems):
        issues.append(
            CheckIssue(
                kind='sheet_file_missing_from_manifest',
                detail=(
                    f'sheets/{stem}.md exists but '
                    f'workbook.toml does not list it'
                ),
                location=f'sheets/{stem}.md',
            )
        )

    issues.extend(_bulk_data_issues(project))

    return issues


_QUOTED_SHEET_RANGE = re.compile(r"^'((?:[^']|'')+)'!")
_BARE_SHEET_RANGE = re.compile(r'^([A-Za-z_][\w]*)!')
_TABLE_REF = re.compile(r'^([A-Za-z_][\w]*)\[')


def _parse_formula_refs(formula: str) -> Tuple[List[str], List[str]]:
    """Return (sheet_refs, name_refs) extracted from `formula`."""
    sheet_refs: List[str] = []
    name_refs: List[str] = []
    for tok in Tokenizer(formula).items:
        if tok.type != 'OPERAND' or tok.subtype != 'RANGE':
            continue
        v = tok.value

        m_q = _QUOTED_SHEET_RANGE.match(v)
        m_b = _BARE_SHEET_RANGE.match(v)
        if m_q:
            sheet_refs.append(m_q.group(1).replace("''", "'"))
            continue
        if m_b:
            sheet_refs.append(m_b.group(1))
            continue

        m_t = _TABLE_REF.match(v)
        if m_t:
            name_refs.append(m_t.group(1))
            continue

        if _is_a1_address(v):
            continue
        if re.fullmatch(r'[A-Za-z_][\w]*', v):
            name_refs.append(v)

    return sheet_refs, name_refs


def _is_a1_address(value: str) -> bool:
    return bool(
        re.fullmatch(r'\$?[A-Za-z]+\$?\d+(?::\$?[A-Za-z]+\$?\d+)?', value)
    )


def _bulk_data_issues(project: Project) -> List[CheckIssue]:
    from ruamel.yaml import YAML

    issues: List[CheckIssue] = []
    yaml = YAML(typ='safe')
    for yp in project.sheets_dir.glob('*.yaml'):
        try:
            doc = yaml.load(yp.read_text()) or {}
        except Exception:
            continue
        src = doc.get('source')
        if isinstance(src, str):
            full = project.root / src
            if not full.is_file():
                issues.append(
                    CheckIssue(
                        kind='bulk_data_missing',
                        detail=(
                            f'{yp.name} declares source={src!r} '
                            f'but file is missing'
                        ),
                        location=str(yp.relative_to(project.root)),
                    )
                )
    return issues
