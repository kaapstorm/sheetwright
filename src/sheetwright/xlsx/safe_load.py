"""Safe openpyxl workbook loader with pre-load security checks."""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
from openpyxl.workbook.workbook import Workbook

from sheetwright.exceptions import XlsxTooLargeError
from sheetwright.security import SecurityLimits

_WORKSHEET_PREFIX = 'xl/worksheets/'
_SHARED_STRINGS = 'xl/sharedStrings.xml'
_RATIO_THRESHOLD = 1000
_SIZE_FLOOR = 8 * 1024 * 1024


def safe_load_workbook(
    path: Path,
    limits: SecurityLimits,
    *,
    data_only: bool = False,
    read_only: bool = False,
) -> Workbook:
    _check_zip(path, limits)
    wb = openpyxl.load_workbook(path, data_only=data_only, read_only=read_only)
    if read_only:
        try:
            _stream_cell_count(wb, limits)
        except BaseException:
            wb.close()
            raise
    return wb


def _check_zip(path: Path, limits: SecurityLimits) -> None:
    with zipfile.ZipFile(path) as zf:
        total_uncompressed = 0
        sheet_count = 0
        for info in zf.infolist():
            total_uncompressed += info.file_size
            if info.filename.startswith(
                _WORKSHEET_PREFIX
            ) and not info.filename.endswith('/'):
                sheet_count += 1
            if info.compress_size:
                ratio = info.file_size / info.compress_size
                if ratio > _RATIO_THRESHOLD and info.file_size > _SIZE_FLOOR:
                    raise XlsxTooLargeError(
                        f'Compression ratio {ratio:.0f}:1 on {info.filename!r} '
                        f'exceeds threshold {_RATIO_THRESHOLD} (file size '
                        f'{info.file_size} > floor {_SIZE_FLOOR}).'
                    )
            if info.filename == _SHARED_STRINGS:
                cap = limits.max_xlsx_shared_strings * 32
                if info.file_size > cap:
                    raise XlsxTooLargeError(
                        f'xl/sharedStrings.xml is {info.file_size:,} bytes '
                        f'uncompressed; cap is {cap:,} '
                        f'(max_xlsx_shared_strings={limits.max_xlsx_shared_strings})'
                    )

        if total_uncompressed > limits.max_xlsx_uncompressed_bytes:
            raise XlsxTooLargeError(
                f'xlsx uncompressed size {total_uncompressed:,} bytes exceeds '
                f'max_xlsx_uncompressed_bytes={limits.max_xlsx_uncompressed_bytes:,}'
            )
        if sheet_count > limits.max_xlsx_sheet_count:
            raise XlsxTooLargeError(
                f'xlsx contains {sheet_count} worksheets; '
                f'max_xlsx_sheet_count={limits.max_xlsx_sheet_count}'
            )


def _stream_cell_count(wb: Workbook, limits: SecurityLimits) -> None:
    for ws in wb.worksheets:
        count = 0
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    count += 1
                    if count > limits.max_xlsx_cells_per_sheet:
                        raise XlsxTooLargeError(
                            f'Sheet {ws.title!r} exceeds '
                            f'max_xlsx_cells_per_sheet={limits.max_xlsx_cells_per_sheet:,}'
                        )
