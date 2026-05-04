"""External-reference detection and flattening.

`detect_external_refs` returns the unique external-ref formula
strings present in an xlsx (empty tuple = none). It uses openpyxl's
formula tokenizer so structured table references (`Table[Col]`) are
not misidentified as external refs.

`flatten_external_refs` mutates a `Workbook` in-place: each cell whose
formula contains an external reference is replaced with a literal
cell holding the cached calculated value (read via
`openpyxl(..., data_only=True)`).
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Tuple

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer

from sheetwright.model.cell import Cell
from sheetwright.model.workbook import Workbook


# An external reference in a tokenized formula has subtype RANGE and
# its value starts with either `'[` (path-quoted) or `[N]` (compiled
# bookId form).
_EXTERNAL_RE = re.compile(r"^(?:'\[|\[\d+\])")


def _is_external_ref_token(value: str) -> bool:
    return bool(_EXTERNAL_RE.match(value))


def detect_external_refs(xlsx_path: Path) -> Tuple[str, ...]:
    """Return unique external-ref tokens present in any formula."""
    src = openpyxl.load_workbook(xlsx_path, data_only=False)
    found: set[str] = set()
    for ws in src.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str) or not v.startswith('='):
                    continue
                for tok in Tokenizer(v).items:
                    if (
                        tok.type == 'OPERAND'
                        and tok.subtype == 'RANGE'
                        and _is_external_ref_token(tok.value)
                    ):
                        found.add(tok.value)
    return tuple(sorted(found))


def flatten_external_refs(wb: Workbook, xlsx_path: Path) -> None:
    """Replace external-ref formulas in `wb` with their cached values.

    `wb` must have been loaded from `xlsx_path` (or an equivalent
    file) — we re-open the same file in data_only mode to pull cached
    values per (sheet, address).
    """
    cached = openpyxl.load_workbook(xlsx_path, data_only=True)
    for sheet in wb.sheets:
        ws = cached[sheet.name]
        for addr, cell in list(sheet.cells.items()):
            if cell.formula is None:
                continue
            if not _formula_has_external_ref(cell.formula):
                continue
            value = ws[addr].value
            sheet.set(
                addr,
                Cell(value=value, format_id=cell.format_id),
            )


def _formula_has_external_ref(formula: str) -> bool:
    for tok in Tokenizer(formula).items:
        if (
            tok.type == 'OPERAND'
            and tok.subtype == 'RANGE'
            and _is_external_ref_token(tok.value)
        ):
            return True
    return False
